"""Generate a presentation summarizing the Conversion Engine from README + RUNBOOK."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "Conversion_Engine_Overview.pptx"

NAVY = RGBColor(0x0E, 0x1F, 0x3A)
ACCENT = RGBColor(0xE8, 0x6A, 0x33)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
TEXT = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5A, 0x6A, 0x85)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_background(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.shadow.inherit = False
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return bg


def add_accent_bar(slide, top=Inches(0.55), height=Inches(0.06)):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), top, Inches(1.2), height)
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT


def add_text(slide, left, top, width, height, text, *, size=18, bold=False, color=TEXT, align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return box


def add_bullets(slide, left, top, width, height, bullets, *, size=16, color=TEXT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"


def add_title(slide, title, subtitle=None):
    add_text(slide, Inches(0.6), Inches(0.25), Inches(12), Inches(0.55),
             title, size=30, bold=True, color=NAVY)
    add_accent_bar(slide)
    if subtitle:
        add_text(slide, Inches(0.6), Inches(0.7), Inches(12), Inches(0.4),
                 subtitle, size=14, color=MUTED)


def add_card(slide, left, top, width, height, title, body, *, accent=ACCENT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.line.color.rgb = RGBColor(0xD9, 0xDF, 0xEA)
    card.line.width = Pt(0.75)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    add_text(slide, left + Inches(0.25), top + Inches(0.12), width - Inches(0.35), Inches(0.4),
             title, size=15, bold=True, color=NAVY)
    add_text(slide, left + Inches(0.25), top + Inches(0.5), width - Inches(0.35), height - Inches(0.55),
             body, size=12, color=TEXT)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ---------- Slide 1: Title ----------
s = prs.slides.add_slide(blank)
add_background(s, NAVY)
hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.8), prs.slide_width, Inches(0.08))
hero.line.fill.background()
hero.fill.solid()
hero.fill.fore_color.rgb = ACCENT
add_text(s, Inches(0.8), Inches(1.5), Inches(12), Inches(0.7),
         "Conversion Engine", size=54, bold=True, color=WHITE)
add_text(s, Inches(0.8), Inches(2.2), Inches(12), Inches(0.5),
         "Automated lead-generation & conversion for Tenacious Consulting", size=22, color=RGBColor(0xC9, 0xD3, 0xE6))
add_text(s, Inches(0.8), Inches(3.1), Inches(12), Inches(0.5),
         "Find the lead. Ground the conversation. Respect the brand. Ship it.",
         size=16, color=ACCENT)
add_text(s, Inches(0.8), Inches(6.6), Inches(12), Inches(0.4),
         "Built from public data only • Kill-switch gated • Trace-grounded",
         size=12, color=RGBColor(0x9A, 0xAA, 0xC4))

# ---------- Slide 2: What's Built ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "What's Built", "Six pillars, one end-to-end pipeline")

cards = [
    ("Signal Enrichment", "Firmographics, hiring velocity, layoffs,\nleadership, AI-maturity, tech stack,\nbench-to-brief match."),
    ("The Agent", "ICP classifier (rule-based, abstains <0.6),\ntone-aware composer (5 markers, 120-word\ncap), reply handler, handoff gate."),
    ("Kill-Switch Gate", "Every outbound action gated by\nTENACIOUS_OUTBOUND_ENABLED.\nDefault routes to staff sinks."),
    ("Channels", "Email (Resend), SMS (Africa's Talking),\nVoice stub. HubSpot CRM + Cal.com\nself-hosted booking."),
    ("Evaluation", "τ²-Bench harness, 34 probes across\n12 categories, failure taxonomy,\nheld-out trace tier."),
    ("Observability", "Langfuse traces & cost. Local fallback\nto data/local_traces.jsonl. Weekly\npolicy audit + final-check gauntlet."),
]
positions = [
    (Inches(0.6), Inches(1.35)),
    (Inches(4.7), Inches(1.35)),
    (Inches(8.8), Inches(1.35)),
    (Inches(0.6), Inches(4.3)),
    (Inches(4.7), Inches(4.3)),
    (Inches(8.8), Inches(4.3)),
]
for (title, body), (l, t) in zip(cards, positions):
    add_card(s, l, t, Inches(3.95), Inches(2.7), title, body)

# ---------- Slide 3: Architecture ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Architecture", "Public-data sources → agent → kill-switch → channels")

def pipeline_box(left, top, w, h, title, sub, fill):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    sh.line.color.rgb = NAVY
    sh.line.width = Pt(1)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    add_text(s, left, top + Inches(0.12), w, Inches(0.35), title,
             size=13, bold=True, color=NAVY, align=2)
    add_text(s, left, top + Inches(0.5), w, h - Inches(0.55), sub,
             size=10, color=TEXT, align=2)

def arrow(x1, y1, x2, y2):
    a = s.shapes.add_connector(1, x1, y1, x2, y2)
    a.line.color.rgb = NAVY
    a.line.width = Pt(1.5)

# Sources row
pipeline_box(Inches(0.6), Inches(1.3), Inches(2.8), Inches(1.2),
             "Public Data", "Crunchbase ODM (frozen)\nlayoffs.fyi · Job posts (≤200/wk)", WHITE)
# Enrichment
pipeline_box(Inches(4.0), Inches(1.3), Inches(5.3), Inches(1.2),
             "Signal-Enrichment Pipeline",
             "firmographics → hiring → layoffs → leadership → AI-maturity → tech stack → bench match",
             WHITE)
# Briefs
pipeline_box(Inches(9.9), Inches(1.3), Inches(2.8), Inches(1.2),
             "Briefs",
             "hiring_signal_brief.json\ncompetitor_gap_brief.json", WHITE)
arrow(Inches(3.4), Inches(1.9), Inches(4.0), Inches(1.9))
arrow(Inches(9.3), Inches(1.9), Inches(9.9), Inches(1.9))

# Agent
pipeline_box(Inches(3.5), Inches(2.8), Inches(6.3), Inches(1.4),
             "The Agent",
             "ICP classifier · Tone-aware composer · Reply handler · Handoff gate",
             RGBColor(0xFF, 0xEC, 0xDD))
arrow(Inches(6.6), Inches(2.5), Inches(6.6), Inches(2.8))

# Kill switch
pipeline_box(Inches(3.5), Inches(4.5), Inches(6.3), Inches(0.9),
             "Kill-Switch Gate  —  TENACIOUS_OUTBOUND_ENABLED?",
             "unset → staff sink     |     set → real sender",
             RGBColor(0xFD, 0xE6, 0xE6))
arrow(Inches(6.6), Inches(4.2), Inches(6.6), Inches(4.5))

# Channels
pipeline_box(Inches(0.6), Inches(5.7), Inches(3.9), Inches(1.0),
             "Email (primary)", "Resend → HubSpot MCP", WHITE)
pipeline_box(Inches(4.7), Inches(5.7), Inches(3.9), Inches(1.0),
             "SMS (scheduling)", "Africa's Talking → Cal.com (self-host)", WHITE)
pipeline_box(Inches(8.8), Inches(5.7), Inches(3.9), Inches(1.0),
             "Voice (bonus)", "Langfuse traces + cost", WHITE)
arrow(Inches(2.5), Inches(5.4), Inches(2.5), Inches(5.7))
arrow(Inches(6.6), Inches(5.4), Inches(6.6), Inches(5.7))
arrow(Inches(10.7), Inches(5.4), Inches(10.7), Inches(5.7))

# ---------- Slide 4: Kill Switch ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Kill Switch — Read This First",
          "TENACIOUS_OUTBOUND_ENABLED gates every outbound action")

# Two-state cards
add_card(s, Inches(0.6), Inches(1.4), Inches(6.0), Inches(2.4),
         "State: Unset (default)",
         "All channels route to staff sinks.\n\nNo real prospect gets contacted.\n\nThis is the safe default for the\nchallenge week.",
         accent=RGBColor(0x2E, 0x8B, 0x57))
add_card(s, Inches(6.8), Inches(1.4), Inches(6.0), Inches(2.4),
         "State: Set to 1",
         "Routes to actual recipients.\n\nRequires explicit staff approval.\n\nRecipient must be sink or in\ndata/synthetic_prospects.json.",
         accent=ACCENT)

# Guards
add_text(s, Inches(0.6), Inches(4.0), Inches(12), Inches(0.4),
         "Additional guards", size=18, bold=True, color=NAVY)
add_bullets(s, Inches(0.8), Inches(4.45), Inches(12), Inches(2.6), [
    "Recipient must be the sink or a synthetic prospect — anything else raises PolicyViolation.",
    "Every email must carry the X-Tenacious-Status: draft header — missing header raises PolicyViolation.",
    "CI grep enforces that only agent/kill_switch.py (and thin adapters) import provider send methods.",
    "No real customer contact during the challenge week. No fabricated numbers. Secrets via .env only.",
], size=14)

# ---------- Slide 5: Quick Start ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Quick Start", "From clone to first send in four moves")

steps = [
    ("1. Setup",
     "make setup\nvi .env       # fill credentials\nmake ack       # policy ack"),
    ("2. Start services",
     "docker compose -f infra/\n  docker-compose.yml up -d\nngrok http 8000\nmake server"),
    ("3. Verify",
     "make smoke\n# 5+ green checks\ncurl localhost:8000/health"),
    ("4. Run",
     "make enrich \\\n  DOMAIN=delamode-group.com\nmake compose-and-send \\\n  DOMAIN=delamode-group.com"),
]
for i, (t, body) in enumerate(steps):
    add_card(s, Inches(0.6 + i * 3.13), Inches(1.4),
             Inches(2.95), Inches(2.6), t, body)

add_text(s, Inches(0.6), Inches(4.3), Inches(12), Inches(0.4),
         "Required free-tier accounts", size=18, bold=True, color=NAVY)
add_bullets(s, Inches(0.8), Inches(4.75), Inches(12), Inches(2.4), [
    "Resend — email delivery (RESEND_API_KEY, RESEND_FROM_ADDRESS)",
    "Africa's Talking — SMS (AT_API_KEY, AT_USERNAME)",
    "HubSpot Dev Sandbox — CRM (HUBSPOT_PRIVATE_APP_TOKEN, HUBSPOT_PORTAL_ID)",
    "Langfuse — observability (LANGFUSE_PUBLIC_KEY, LANGFUSE_SECRET_KEY)",
    "OpenRouter — LLM access (OPENROUTER_API_KEY)",
], size=13)

# ---------- Slide 6: Runbook — Four Terminals ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Runbook — Four Terminals", "Long-lived processes plus one for ad-hoc runs")

rows = [
    ("A", "Cal.com + Postgres",
     "docker compose -f infra/docker-compose.yml up -d", "long-lived"),
    ("B", "Webhook tunnel",
     "ngrok http 8000   (or cloudflared tunnel)", "long-lived"),
    ("C", "FastAPI server",
     "make server   # uvicorn on :8000", "long-lived"),
    ("D", "Ad-hoc runs",
     "make enrich … / make compose-and-send …", "short-lived"),
]
header_y = Inches(1.35)
header = ["Term", "Process", "Command", "Lifetime"]
widths = [Inches(0.9), Inches(2.8), Inches(7.4), Inches(1.6)]
left = Inches(0.6)
x = left
for h, w in zip(header, widths):
    cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, header_y, w, Inches(0.5))
    cell.line.color.rgb = NAVY
    cell.fill.solid()
    cell.fill.fore_color.rgb = NAVY
    add_text(s, x + Inches(0.1), header_y + Inches(0.1), w, Inches(0.3),
             h, size=13, bold=True, color=WHITE)
    x += w
y = header_y + Inches(0.5)
for row in rows:
    x = left
    for v, w in zip(row, widths):
        cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.55))
        cell.line.color.rgb = RGBColor(0xCC, 0xD3, 0xE0)
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
        add_text(s, x + Inches(0.1), y + Inches(0.13), w, Inches(0.35),
                 v, size=12, color=TEXT)
        x += w
    y += Inches(0.55)

add_text(s, Inches(0.6), Inches(4.7), Inches(12), Inches(0.4),
         "Note", size=15, bold=True, color=NAVY)
add_bullets(s, Inches(0.8), Inches(5.1), Inches(12), Inches(2), [
    "HubSpot MCP is not a separate daemon — the Python agent spawns it over stdio per-process.",
    "Langfuse is cloud — no daemon. With keys absent, traces fall back to data/local_traces.jsonl.",
    "First Cal.com boot: create admin + two event types (discovery-15, discovery-30); paste API key into .env.",
], size=13)

# ---------- Slide 7: Operating the System ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Operating the System", "Day-to-day commands")

groups = [
    ("Enrichment & outbound",
     "make enrich DOMAIN=<domain>\nmake compose-and-send DOMAIN=<domain>"),
    ("τ²-Bench evaluation",
     "make tau2-baseline\nTAU2_SEALED_ACCESS=1 EVAL_TIER_ENABLED=1 \\\n  make tau2-eval"),
    ("Probes & memo",
     "make probes\nmake probes P=P-0001\nmake memo"),
    ("Ops",
     "make server  /  make test  /  make lint\nmake audit  /  make final-check"),
]
positions = [
    (Inches(0.6), Inches(1.4)),
    (Inches(6.95), Inches(1.4)),
    (Inches(0.6), Inches(4.3)),
    (Inches(6.95), Inches(4.3)),
]
for (t, body), (l, top) in zip(groups, positions):
    add_card(s, l, top, Inches(5.85), Inches(2.6), t, body)

# ---------- Slide 8: Demo Script ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Demo Script", "≤8 minutes, ten beats — proves the system end-to-end")

beats = [
    ("Step 1 — Enrichment Live (1:00)",
     "make enrich DOMAIN=delamode-group.com — Langfuse spans, briefs written, per-signal confidence visible."),
    ("Step 2 — Cold Email Compose & Send (1:00)",
     "make compose-and-send … — segment-specific draft, tone-check passes, sink: True, HubSpot record populates."),
    ("Step 3 — Engaged Reply (1:00)",
     "Reply from sink inbox → classifier tags engaged → warm response with Cal.com link → HubSpot updated."),
    ("Step 4 — SMS Scheduling Handoff (0:45)",
     "Phone shared → channel switches to SMS → SMS sink confirms → Langfuse cross-channel trace."),
    ("Step 5 — Cal.com Booking (0:45)",
     "Slot picked → booking written to data/calcom_local/bookings.jsonl → context brief NOTE on HubSpot Deal."),
    ("Step 6 — Abstention Path (0:30)",
     "make compose-and-send DOMAIN=windowclassics.com — agent refuses to assert aggressive hiring; softer language."),
    ("Step 7 — Classification Nuance (0:30)",
     "Post-layoff + funding prospect → Segment 2 not Segment 1; rule visible in Langfuse."),
    ("Step 8 — τ²-Bench Score (0:30)",
     "make tau2-baseline — harness produces a trace; pass@1 visible."),
    ("Step 9 — Probe Walkthrough (0:45)",
     "make probes P=P-0001 — show probe definition + before/after trigger rate."),
    ("Step 10 — Outro (0:20)",
     "Pilot recommendation + kill-switch clause: trigger metric, threshold, rollback condition."),
]
y = Inches(1.3)
for i, (t, body) in enumerate(beats):
    col = i % 2
    row = i // 2
    left = Inches(0.6 + col * 6.3)
    top = Inches(1.3 + row * 1.18)
    add_card(s, left, top, Inches(6.1), Inches(1.05), t, body)

# ---------- Slide 9: Limitations & Handoff ----------
s = prs.slides.add_slide(blank)
add_background(s, LIGHT)
add_title(s, "Limitations & Handoff to Humans",
          "Five hard rails escalate before the brand or prospect can be hurt")

triggers = [
    ("Regulatory keyword in reply",
     "MSA / DPA / BAA / SOW / NDA / DPIA / GDPR / HIPAA → hard handoff; agent stops sending."),
    ("Pricing question outside published bands",
     "Hard handoff. Agent never quotes a number it cannot ground."),
    ("Bench over-commit",
     "Prospect asks for stack with bench count = 0 → hard handoff; one acknowledgement, no warm draft."),
    ("Tone-check double-fail",
     "Soft handoff; logs to eval/runs/tone_flagged.jsonl."),
    ("Discovery call booked",
     "Soft handoff; context brief attached to HubSpot Deal as a NOTE."),
]
y = Inches(1.35)
for t, body in triggers:
    add_card(s, Inches(0.6), y, Inches(12.13), Inches(0.95), t, body)
    y += Inches(1.05)

# ---------- Slide 10: Non-Negotiables / Closing ----------
s = prs.slides.add_slide(blank)
add_background(s, NAVY)
add_text(s, Inches(0.8), Inches(0.6), Inches(12), Inches(0.7),
         "Non-Negotiables", size=36, bold=True, color=WHITE)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(1.2), Inches(0.06))
bar.line.fill.background()
bar.fill.solid()
bar.fill.fore_color.rgb = ACCENT

items = [
    "Kill switch defaults unset.",
    "No real customer contact during the challenge week.",
    "Every outbound carries X-Tenacious-Status: draft.",
    "No hard-coded secrets — everything via .env or config.yaml.",
    "No fabricated Tenacious numbers — every claim resolves to a trace, seed row, or public source.",
]
y = Inches(1.7)
for it in items:
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), y + Inches(0.12), Inches(0.18), Inches(0.18))
    dot.line.fill.background()
    dot.fill.solid()
    dot.fill.fore_color.rgb = ACCENT
    add_text(s, Inches(1.3), y, Inches(11.5), Inches(0.5),
             it, size=18, color=WHITE)
    y += Inches(0.7)

add_text(s, Inches(0.8), Inches(6.5), Inches(12), Inches(0.4),
         "Specs · Plans · RUNBOOK.md  —  ship it.",
         size=14, color=ACCENT)

prs.save(OUT)
print(f"Wrote {OUT}")
